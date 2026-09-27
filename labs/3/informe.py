import json
import os

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt, RGBColor

RESULTADOS = "resultados"
CAPTURAS = "capturas"
SALIDA = "Informe_Laboratorio03.docx"
LOGO = "utec.png"


def metricas():
    with open(os.path.join(RESULTADOS, "metricas.json")) as fuente:
        filas = json.load(fuente)
    for fila in filas:
        for algoritmo in ("sort", "hash"):
            datos = fila[algoritmo]
            datos["io_total"] = datos["pages_read"] + datos["pages_written"]
    return filas


def coma(valor, decimales=2):
    return f"{valor:.{decimales}f}".replace(".", ",")


def mil(valor):
    return f"{int(valor):,}".replace(",", " ")


def documento():
    doc = Document()
    seccion = doc.sections[0]
    seccion.page_width = Cm(21.0)
    seccion.page_height = Cm(29.7)
    seccion.top_margin = Cm(2.5)
    seccion.bottom_margin = Cm(2.5)
    seccion.left_margin = Cm(3.0)
    seccion.right_margin = Cm(3.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for nombre, tamano in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
        estilo = doc.styles[nombre]
        estilo.font.name = "Arial"
        estilo.font.size = Pt(tamano)
        estilo.font.bold = True
        estilo.font.italic = nombre == "Heading 3"
        estilo.font.color.rgb = RGBColor(0, 0, 0)
        estilo.paragraph_format.line_spacing = 1.5
        estilo.paragraph_format.space_before = Pt(12)
        estilo.paragraph_format.space_after = Pt(6)
        estilo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    pie = doc.styles.add_style("Pie", WD_STYLE_TYPE.PARAGRAPH)
    pie.font.name = "Arial"
    pie.font.size = Pt(10)
    pie.font.italic = True
    pie.paragraph_format.line_spacing = 1.0
    pie.paragraph_format.space_after = Pt(10)
    pie.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return doc


def parrafo(doc, texto, estilo=None):
    return doc.add_paragraph(texto, style=estilo)


def captura(doc, nombre, pie):
    figura(doc, os.path.join(CAPTURAS, nombre), pie, ancho=15.0)


def figura(doc, ruta, pie, ancho=14.0):
    doc.add_picture(ruta, width=Cm(ancho))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(pie, style="Pie")


def tabla(doc, cabeceras, filas):
    tabla = doc.add_table(rows=1, cols=len(cabeceras))
    tabla.style = "Table Grid"
    for celda, texto in zip(tabla.rows[0].cells, cabeceras):
        celda.text = texto
        for parrafo_celda in celda.paragraphs:
            parrafo_celda.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for corrida in parrafo_celda.runs:
                corrida.font.bold = True
    for fila in filas:
        celdas = tabla.add_row().cells
        for indice, (celda, texto) in enumerate(zip(celdas, fila)):
            celda.text = str(texto)
            for parrafo_celda in celda.paragraphs:
                parrafo_celda.alignment = (WD_ALIGN_PARAGRAPH.LEFT if indice == 0
                                           else WD_ALIGN_PARAGRAPH.CENTER)
    for fila in tabla.rows:
        for celda in fila.cells:
            for parrafo_celda in celda.paragraphs:
                parrafo_celda.paragraph_format.line_spacing = 1.0
                parrafo_celda.paragraph_format.space_after = Pt(2)
                for corrida in parrafo_celda.runs:
                    corrida.font.name = "Arial"
                    corrida.font.size = Pt(9)
    doc.add_paragraph()


def portada(doc):
    def centrado(texto, tamano=12, negrita=False, espacio=6):
        parrafo_actual = doc.add_paragraph()
        parrafo_actual.alignment = WD_ALIGN_PARAGRAPH.CENTER
        parrafo_actual.paragraph_format.space_after = Pt(espacio)
        corrida = parrafo_actual.add_run(texto)
        corrida.font.size = Pt(tamano)
        corrida.font.bold = negrita
        return parrafo_actual

    centrado("UNIVERSIDAD DE INGENIERÍA Y TECNOLOGÍA", 14, True, 24)
    doc.add_picture(LOGO, width=Cm(6))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(24)
    centrado("Curso:", 14, True, 2)
    centrado("Base de Datos II", 14, False, 36)
    centrado("Laboratorio 03: External Algorithms", 18, True, 48)
    centrado("Alumno:", 12, True, 2)
    centrado("Llerena Silva, Nicolás Alejandro (202110190)", 12, False, 24)
    centrado("Profesor:", 12, True, 2)
    centrado("Lovon, Percy", 12, False, 36)
    centrado("Lima, agosto de 2026", 12, False, 0)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def parte1(doc):
    doc.add_heading("Introducción", level=1)
    parrafo(doc, "Este laboratorio tiene dos mitades. En la primera reviso qué algoritmos de memoria "
                 "externa usa PostgreSQL cuando lo obligo a trabajar con 64 kB de work_mem sobre la base "
                 "employees del Laboratorio 01. En la segunda implemento en Python, sobre archivos "
                 "binarios paginados, el Two-Phase Multiway Merge Sort y el External Hashing para un "
                 "GROUP BY, y contrasto los resultados contra los del motor.")
    parrafo(doc, "Las pruebas se corrieron sobre PostgreSQL 14.18 en macOS. Las tablas involucradas son "
                 "employee, con 300 024 filas, y department_employee, con 331 603. En la implementación "
                 "en Python el tamaño de página es de 4096 bytes. Las salidas que aparecen a lo largo del "
                 "informe son las que devolvieron psql y los programas; los archivos completos están en la "
                 "carpeta evidencia del entregable.")

    doc.add_heading("Parte 1: planes de ejecución en PostgreSQL", level=1)
    parrafo(doc, "Cada consulta se ejecutó dos veces, primero con work_mem en 64 kB y después en 2 MB, "
                 "dentro de una transacción y con EXPLAIN (ANALYZE, BUFFERS). El script "
                 "capturar_planes.sh automatiza la captura y deja cada plan en un archivo aparte. Como "
                 "EXPLAIN no informa cuántas pasadas de merge hizo un ordenamiento, para la primera "
                 "consulta también activé trace_sort junto con client_min_messages = LOG, que es lo que "
                 "hace que esa traza interna llegue al cliente, y apagué el paralelismo para poder seguir "
                 "un solo ordenamiento.")

    doc.add_heading("Consulta 1: ORDER BY sobre hire_date", level=2)
    captura(doc, "cap_q1_64kb.png", "Plan de la consulta 1 con work_mem = 64kB.")
    parrafo(doc, "El ordenamiento no cabe en memoria. El nodo Sort reporta el método external merge y "
                 "vuelca 7192 kB a disco, y el worker que lo acompaña otros 6632 kB; entre ambos el motor "
                 "leyó 6937 bloques temporales y escribió 7519. Al subir work_mem a 2 MB el plan mantiene "
                 "exactamente la misma forma y el método sigue siendo externo, porque las 300 024 filas "
                 "ocupan alrededor de 7 MB una vez proyectadas. Lo que cambia es el trabajo temporal, que "
                 "baja a 2288 bloques leídos y 2302 escritos, y el tiempo, que pasa de 115,7 ms a 55,0 ms. "
                 "Para ver el caso en memoria hay que llegar a 64 MB: recién ahí el método cambia a "
                 "quicksort con 35 729 kB y desaparecen los archivos temporales.")
    captura(doc, "cap_q1_2mb.png", "El mismo plan con work_mem = 2MB.")
    captura(doc, "cap_q1_64mb.png", "Con 64 MB el ordenamiento ya cabe en memoria.")
    captura(doc, "cap_q1_trace.png", "Conteo de runs y merge steps sobre la traza de trace_sort.")

    parrafo(doc, "¿Qué algoritmo de memoria externa aparece cuando la RAM es insuficiente? El external "
                 "merge sort, que es la versión multipasada del Two-Phase Multiway Merge Sort. PostgreSQL "
                 "llena work_mem, ordena ese bloque con quicksort y lo escribe como un run en una cinta "
                 "temporal, que es la fase 1; después mezcla las cintas con un árbol de torneo, que cumple "
                 "el papel del min-heap de la fase 2.")
    parrafo(doc, "¿Cuántos merge passes realizó? La traza con 64 kB muestra 586 runs y 118 merge steps "
                 "repartidos sobre 7 cintas, seis de entrada y una de salida. Eso equivale a cuatro "
                 "pasadas: 586 entre 6 dan unas 98 mezclas en el primer nivel, después 17, luego 3 y una "
                 "final, que suman 119 y coinciden con las 118 observadas. Con 2 MB los runs bajan a 16 y "
                 "alcanzan 2 merge steps. La relación con la fase 1 es directa: si hay F = B − 1 cintas de "
                 "entrada, el número de pasadas es el logaritmo en base F de la cantidad de runs, y solo "
                 "cuando esos runs caben en las cintas disponibles el algoritmo es de dos fases en el "
                 "sentido estricto. Más memoria produce runs más largos, menos runs y por lo tanto menos "
                 "pasadas.")
    parrafo(doc, "¿Cómo cambia el costo al aumentar work_mem? El costo estimado del plan baja de 55 401 a "
                 "45 748 entre 64 kB y 2 MB, y el tiempo real se reduce a la mitad, aunque el método siga "
                 "siendo externo: lo que mejora es el número de pasadas, y eso se ve en el I/O temporal, "
                 "que cae más de tres veces. Con 64 MB el costo total llega a 33 526, el tiempo a 46,9 ms "
                 "y el planificador incluso abandona el plan paralelo, porque ya no necesita repartir el "
                 "ordenamiento entre dos procesos.")

    doc.add_heading("Consulta 2: GROUP BY sobre from_date", level=2)
    captura(doc, "cap_q2_64kb.png", "Plan de la consulta 2 con work_mem = 64kB.")
    parrafo(doc, "Este plan es interesante porque con poca memoria aparecen los dos algoritmos a la vez. "
                 "La agregación parcial de cada proceso se resuelve con un Partial HashAggregate que "
                 "desborda a disco, y los resultados parciales se combinan con Sort más Finalize "
                 "GroupAggregate, porque Gather Merge exige entrada ordenada. Con 2 MB los 6393 grupos "
                 "entran en 721 kB, el nodo pasa a Batches: 1 y el Sort desaparece del plan.")
    captura(doc, "cap_q2_2mb.png", "El mismo plan con work_mem = 2MB.")
    parrafo(doc, "¿Hash Aggregate o Sort + Group Aggregate? Las dos. Desde la versión 13 PostgreSQL ya no "
                 "descarta el hash aggregate cuando no cabe en memoria, sino que lo particiona: el plan "
                 "muestra Planned Partitions: 4, Batches: 21, 93 kB de memoria y 6920 kB en disco. El Sort "
                 "que aparece encima no es una alternativa al hash, sino el paso que ordena los resultados "
                 "parciales antes de combinarlos.")
    parrafo(doc, "Cuando aparece Sort + Group Aggregate, internamente se está usando el mismo external "
                 "merge sort de la consulta anterior; aquí el nodo reporta external merge con 184 kB en "
                 "disco. Es la agregación basada en ordenamiento: se ordena por la clave de grupo y una "
                 "sola pasada secuencial acumula los conteos de cada corrida de claves iguales.")
    parrafo(doc, "Cuando aparece Hash Aggregate con batches mayores que uno, lo que está ocurriendo es "
                 "external hashing, en su variante grace o híbrida. Es exactamente lo que implementé en "
                 "external_hashing.py: la fase 1 reparte las tuplas en particiones según una función hash "
                 "sobre from_date y las escribe en archivos temporales, y la fase 2 procesa una partición "
                 "a la vez construyendo la tabla hash en RAM. El número de batches es el número de "
                 "particiones, y crece cuando la tabla hash no cabe en work_mem.")

    doc.add_heading("Consulta 3: JOIN sobre employee_id", level=2)
    captura(doc, "cap_q3_64kb.png", "Plan de la consulta 3 con work_mem = 64kB.")
    parrafo(doc, "El algoritmo elegido con RAM limitada es un Hash Join en su variante externa, el grace "
                 "hash join. La tabla de build es employee y no entra en 64 kB, así que el nodo Hash la "
                 "parte en 512 batches escritos a disco y la relación de probe se reparte con la misma "
                 "función hash, lo que suma 4422 bloques temporales leídos y otros tantos escritos. Con "
                 "2 MB el mismo plan usa 16 batches y 1526 kB de memoria, y el tiempo cae de 537,8 ms a "
                 "110,8 ms.")
    captura(doc, "cap_q3_2mb.png", "El mismo plan con work_mem = 2MB.")
    parrafo(doc, "Esos batches son particiones, y se corresponden uno a uno con la fase 1 del external "
                 "hashing: PostgreSQL elige la potencia de dos más pequeña tal que un batch de la relación "
                 "de build quepa en work_mem, y como reparte ambas relaciones con la misma función, las "
                 "tuplas que hacen match caen siempre en el mismo batch. Eso permite que la fase 2 procese "
                 "cada par de particiones de forma independiente en memoria. Pasar de 16 a 512 particiones "
                 "multiplicó el tiempo casi por cinco por el costo de escribir y releer tantos archivos "
                 "temporales.")
    parrafo(doc, "Merge Join no apareció en ninguno de los dos escenarios. Si el planificador lo hubiera "
                 "elegido, habría necesitado dejar ambas entradas ordenadas por employee_id antes del "
                 "join, sea con dos nodos Sort, con el gasto de disco que eso implica con 64 kB, o "
                 "recorriendo los índices de clave primaria, que ya entregan las filas en ese orden.")


def parte2(doc, filas):
    s64 = next(f["sort"] for f in filas if f["buffer_size"] == 65536)
    s512 = next(f["sort"] for f in filas if f["buffer_size"] == 524288)
    h64 = next(f["hash"] for f in filas if f["buffer_size"] == 65536)
    h1024 = next(f["hash"] for f in filas if f["buffer_size"] == 1048576)

    doc.add_heading("Parte 2: implementación en Python", level=1)

    doc.add_heading("Estructura del heap file", level=2)
    parrafo(doc, "Los dos archivos binarios se generan desde PostgreSQL con export_data.py, que exporta "
                 "cada tabla a CSV y la vuelca al formato paginado. La página 0 guarda la cabecera del "
                 "archivo: número mágico, tamaño de página, tamaño de registro, cantidad de páginas y de "
                 "registros, la cadena de formato de struct y los nombres de los campos. Guardar el "
                 "formato dentro del archivo es lo que permite que read_page(heap_path, page_id, "
                 "page_size) devuelva tuplas ya tipadas sin recibir el formato como parámetro, tal como "
                 "pide la firma del enunciado.")
    parrafo(doc, "Cada página de datos empieza con cuatro bytes que indican cuántos registros contiene y "
                 "sigue con registros de longitud fija que nunca cruzan el límite de la página. Para "
                 "employee el registro ocupa 59 bytes y entran 69 por página, de modo que la tabla ocupa "
                 "4349 páginas; para department_employee son 32 bytes, 127 registros por página y 2612 "
                 "páginas. Las cuatro funciones que pide el enunciado son la interfaz pública, pero por "
                 "dentro se apoyan en una clase HeapFile que mantiene el descriptor abierto y lleva la "
                 "cuenta de páginas leídas y escritas. Ese contador es el que alimenta las métricas de "
                 "I/O, y como los runs y las particiones temporales también son heap files, todo el I/O "
                 "queda medido con el mismo instrumento.")
    captura(doc, "cap_heap.png", "Salida de heap_file.py sobre los dos archivos exportados.")

    doc.add_heading("External sorting: Two-Phase Multiway Merge Sort", level=2)
    parrafo(doc, "La fase 1, en generate_runs, lee exactamente B = BUFFER_SIZE // PAGE_SIZE páginas, las "
                 "ordena en memoria por la clave y las escribe como un run. El número de runs es el techo "
                 "de P entre B, y el programa lo verifica con un assert. La fase 2, en multiway_merge, "
                 "mezcla los runs con un min-heap de heapq usando B − 1 buffers de entrada y uno de "
                 "salida: cada lector mantiene una sola página en memoria y pide la siguiente cuando la "
                 "agota, y el buffer de salida se vacía a disco apenas se llena, así que en ningún momento "
                 "hay más de B páginas residentes.")
    parrafo(doc, f"Con este volumen de datos el número de runs supera a los buffers de entrada disponibles "
                 f"en las configuraciones chicas: con 64 KB salen {s64['runs_generated']} runs contra "
                 f"{s64['fan_in']} buffers, así que una sola pasada de mezcla es imposible y la fase 2 se "
                 f"generaliza a varias pasadas, igual que hace PostgreSQL con sus cintas. Desde 512 KB el "
                 f"algoritmo sí es estrictamente de dos fases, porque los {s512['runs_generated']} runs "
                 f"caben en los {s512['fan_in']} buffers, y ahí el I/O toca su mínimo teórico de 4P "
                 f"páginas. La salida se valida recorriéndola entera: orden no decreciente y los 300 024 "
                 f"registros completos.")
    captura(doc, "cap_sort.png", "Salida de external_sort.py con BUFFER_SIZE = 64 KB.")

    doc.add_heading("External hashing para el GROUP BY", level=2)
    parrafo(doc, "La fase 1, en partition_data, recorre el heap file página a página con un solo buffer de "
                 "entrada y reparte cada tupla en k = B − 1 particiones según h_p(from_date) módulo k, con "
                 "un buffer de salida de una página por partición. La fase 2, en aggregate_partitions, "
                 "procesa una partición a la vez y construye una tabla hash residente con una función "
                 "distinta, h_r, implementada con encadenamiento y duplicación de buckets cuando el factor "
                 "de carga pasa de dos. Las dos funciones son crc32 con semillas diferentes, lo que las "
                 "hace deterministas entre ejecuciones; hash() de Python no sirve para esto porque está "
                 "aleatorizado por proceso.")
    parrafo(doc, f"Como las particiones son disjuntas por clave, cada valor de from_date aparece en una "
                 f"sola partición y los conteos se unen sin combinar parciales. Vale la pena notar que la "
                 f"memoria residente no depende del tamaño de la partición sino de sus grupos distintos: "
                 f"con 64 KB la tabla más grande tuvo {h64['max_groups_in_partition']} entradas, unos 7 KB "
                 f"entre claves y contadores, muy por debajo del presupuesto, así que no hizo falta "
                 f"particionar de forma recursiva. El resultado se compara contra el archivo "
                 f"group_by_referencia.csv, que es la misma consulta ejecutada en PostgreSQL, y la "
                 f"comparación es un assert dentro del programa: los 6393 grupos y los 331 603 conteos "
                 f"coinciden exactamente.")
    captura(doc, "cap_hash.png",
            "Salida de external_hashing.py con BUFFER_SIZE = 64 KB. La última línea es la verificación "
            "contra PostgreSQL.")

    doc.add_heading("Análisis de rendimiento", level=2)
    parrafo(doc, "Cada configuración se ejecutó con benchmark.py, que además vuelve a validar el orden de "
                 "la salida y la igualdad del GROUP BY. Las tres primeras filas son las que pide el "
                 "enunciado; agregué 512 KB y 1 MB para mostrar el punto en que el sort se vuelve "
                 "estrictamente de dos fases.")
    captura(doc, "cap_bench.png", "Salida de benchmark.py con las cinco configuraciones.")
    tabla(doc,
          ["BUFFER_SIZE", "B", "Runs", "Pasadas", "Fase 1 (s)", "Fase 2 (s)", "Total (s)", "I/O (págs.)"],
          [[f"{fila['buffer_size'] // 1024} KB", fila["sort"]["buffer_pages"],
            fila["sort"]["runs_generated"], fila["sort"]["merge_passes"],
            coma(fila["sort"]["time_phase1_sec"]), coma(fila["sort"]["time_phase2_sec"]),
            coma(fila["sort"]["time_total_sec"]), mil(fila["sort"]["io_total"])] for fila in filas])
    doc.add_paragraph("External sorting sobre employee, ordenando por hire_date (P = 4349 páginas).",
                      style="Pie")
    tabla(doc,
          ["BUFFER_SIZE", "B", "Particiones", "Grupos máx.", "Fase 1 (s)", "Fase 2 (s)", "Total (s)",
           "I/O (págs.)"],
          [[f"{fila['buffer_size'] // 1024} KB", fila["hash"]["buffer_pages"],
            fila["hash"]["partitions_created"], fila["hash"]["max_groups_in_partition"],
            coma(fila["hash"]["time_phase1_sec"]), coma(fila["hash"]["time_phase2_sec"]),
            coma(fila["hash"]["time_total_sec"]), mil(fila["hash"]["io_total"])] for fila in filas])
    doc.add_paragraph("External hashing sobre department_employee, agrupando por from_date "
                      "(P = 2612 páginas).", style="Pie")
    figura(doc, os.path.join(RESULTADOS, "tiempo_total.png"),
           "Tiempo total contra BUFFER_SIZE para los dos algoritmos.")
    figura(doc, os.path.join(RESULTADOS, "io_total.png"),
           "Páginas leídas más escritas contra BUFFER_SIZE.")

    parrafo(doc, "El número de runs de la fase 1 se reduce a la mitad cada vez que se duplica "
                 "BUFFER_SIZE, porque es el techo de P entre B y B crece de forma proporcional: 272, 136, "
                 "68, 34 y 17. El tiempo, en cambio, no se reduce a la mitad. La fase 1 siempre lee y "
                 "escribe las mismas 2P páginas y su costo se mantiene alrededor de medio segundo; lo que "
                 "mejora es la fase 2, cuyas pasadas caen de tres a dos y después a una. Por eso 128 KB y "
                 "256 KB rinden casi igual, ambas con dos pasadas, y el salto real ocurre en 512 KB.")
    parrafo(doc, f"En el external hashing el número de particiones casi no mueve el costo de I/O, y lo "
                 f"mueve en sentido contrario. El costo es fijo en unas 3P páginas: se lee la relación una "
                 f"vez, se escriben las particiones una vez y se vuelven a leer en la fase 2, "
                 f"{mil(h64['io_total'])} páginas con 15 particiones. Al subir k el I/O crece apenas, "
                 f"hasta {mil(h1024['io_total'])} páginas con 255, y la razón es que cada partición "
                 f"desperdicia en promedio media página final. El riesgo verdadero está en el otro "
                 f"extremo: si k es demasiado chico, la tabla hash de alguna partición no cabe en RAM y "
                 f"hay que particionar de nuevo, y cada nivel extra cuesta otras 2P páginas.")
    parrafo(doc, f"Comparado con PostgreSQL, mi implementación es entre ocho y diecinueve veces más lenta: "
                 f"el motor ordenó en 115,7 ms y agrupó en 44,1 ms, mientras que Python necesitó "
                 f"{coma(s64['time_total_sec'])} s y {coma(h64['time_total_sec'])} s con el mismo "
                 f"presupuesto de memoria. La diferencia no está en el I/O, que es del mismo orden en "
                 f"ambos casos: PostgreSQL movió unos 112 MB de archivos temporales en el ORDER BY y mi "
                 f"sort unos 135 MB. Está en la CPU y en el formato. El motor está escrito en C, compara "
                 f"fechas como enteros de cuatro bytes y mantiene las tuplas en su representación binaria, "
                 f"mientras que Python vuelve a interpretar cada tupla con struct y decodifica seis "
                 f"cadenas por registro en cada pasada. A eso se suma que PostgreSQL repartió el trabajo "
                 f"entre dos procesos y leyó las tablas desde el buffer pool ya caliente, sin I/O físico. "
                 f"Proyectar solo las columnas necesarias antes de ordenar o particionar, como hace el "
                 f"motor, sería la optimización más rentable de mi versión.")
    parrafo(doc, "Entre los dos algoritmos, preferiría external sorting cuando el orden importa o se "
                 "reutiliza más adelante: un ORDER BY, un DISTINCT que deba salir ordenado, la "
                 "construcción de un índice B+ o un merge join posterior. También cuando hay sesgo en los "
                 "datos, porque su costo no depende de cómo se distribuyan las claves, y cuando la "
                 "cantidad de grupos distintos se acerca al número de tuplas, caso en el que la tabla hash "
                 "dejaría de caber. External hashing conviene cuando basta la igualdad y los grupos son "
                 "pocos frente a las tuplas, que es justo lo que pasa acá: 331 603 tuplas se resumen en "
                 "6393 grupos, el hashing cuesta unas 3P páginas contra 8P del sort y termina en menos de "
                 "un tercio del tiempo. Su punto débil es la sensibilidad al sesgo, ya que una clave muy "
                 "frecuente desbalancea una partición, y que no sirve para rangos, desigualdades ni para "
                 "producir salida ordenada.")

    doc.add_heading("Conclusiones", level=1)
    parrafo(doc, "Los tres planes de la primera parte muestran los mismos algoritmos que se implementaron "
                 "en la segunda: external merge sort en el ORDER BY y en el Sort del GROUP BY, y external "
                 "hashing tanto en el HashAggregate desbordado como en los 512 batches del Hash Join. La "
                 "implementación en Python reproduce ese comportamiento con las mismas restricciones de "
                 "memoria y, en el caso del GROUP BY, llega exactamente al mismo resultado que el motor. "
                 "Lo que más se nota al variar BUFFER_SIZE es que la memoria solo ayuda al sort cuando "
                 "alcanza para eliminar una pasada de merge, mientras que el hashing se mantiene plano "
                 "porque su costo de I/O ya es el mínimo posible para el problema.")

    doc.add_heading("Anexo: archivos entregados", level=1)
    parrafo(doc, "Junto con este informe van los tres módulos que pide el enunciado. heap_file.py "
                 "implementa export_to_heap, read_page, write_page y count_pages sobre la clase HeapFile. "
                 "external_sort.py implementa generate_runs, multiway_merge y external_sort, más la "
                 "verificación de orden de la salida. external_hashing.py implementa partition_data, "
                 "aggregate_partitions y external_hash_group_by, junto con las funciones h_p y h_r y la "
                 "tabla hash con encadenamiento. Los datos van en data/employee.bin y "
                 "data/department_employee.bin, exportados desde PostgreSQL con el formato paginado que "
                 "describe la sección 2.1. Cada módulo se ejecuta por sí solo y vuelve a imprimir las "
                 "salidas que aparecen en las capturas de este informe.")


def main():
    filas = metricas()
    doc = documento()
    portada(doc)
    parte1(doc)
    parte2(doc, filas)
    doc.save(SALIDA)
    print(f"{SALIDA} ({os.path.getsize(SALIDA) // 1024} KB)")


if __name__ == "__main__":
    main()
