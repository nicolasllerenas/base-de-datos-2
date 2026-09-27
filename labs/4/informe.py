import json
import os
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt, RGBColor

RESULTADOS = "resultados"
CAPTURAS = "capturas"
SALIDA = "Informe_Laboratorio04.docx"
LOGO = "utec.png"


def metricas():
    with open(os.path.join(RESULTADOS, "metricas.json")) as fuente:
        return json.load(fuente)


def por_tamano(filas, tamano):
    return next(fila for fila in filas if fila["tamano"] == tamano)


def coma(valor, decimales=2):
    return f"{valor:.{decimales}f}".replace(".", ",")


def mil(valor):
    return f"{int(valor):,}".replace(",", " ")


def documento():
    doc = Document()
    seccion = doc.sections[0]
    seccion.page_width = Cm(21.0)
    seccion.page_height = Cm(29.7)
    seccion.top_margin = Cm(2.5)
    seccion.bottom_margin = Cm(2.5)
    seccion.left_margin = Cm(3.0)
    seccion.right_margin = Cm(3.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for nombre, tamano in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
        estilo = doc.styles[nombre]
        estilo.font.name = "Arial"
        estilo.font.size = Pt(tamano)
        estilo.font.bold = True
        estilo.font.italic = nombre == "Heading 3"
        estilo.font.color.rgb = RGBColor(0, 0, 0)
        estilo.paragraph_format.line_spacing = 1.5
        estilo.paragraph_format.space_before = Pt(12)
        estilo.paragraph_format.space_after = Pt(6)
        estilo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    pie = doc.styles.add_style("Pie", WD_STYLE_TYPE.PARAGRAPH)
    pie.font.name = "Arial"
    pie.font.size = Pt(10)
    pie.font.italic = True
    pie.paragraph_format.line_spacing = 1.0
    pie.paragraph_format.space_after = Pt(10)
    pie.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return doc


def parrafo(doc, texto, estilo=None):
    parrafo_actual = doc.add_paragraph(style=estilo)
    for indice, fragmento in enumerate(re.split(r"</?b>", texto)):
        if fragmento:
            parrafo_actual.add_run(fragmento).font.bold = indice % 2 == 1
    return parrafo_actual


def captura(doc, nombre, pie):
    figura(doc, os.path.join(CAPTURAS, nombre), pie, ancho=15.0)


def figura(doc, ruta, pie, ancho=14.0):
    doc.add_picture(ruta, width=Cm(ancho))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(pie, style="Pie")


def tabla(doc, cabeceras, filas, izquierda=(0,)):
    tabla = doc.add_table(rows=1, cols=len(cabeceras))
    tabla.style = "Table Grid"
    for celda, texto in zip(tabla.rows[0].cells, cabeceras):
        celda.text = texto
        for parrafo_celda in celda.paragraphs:
            parrafo_celda.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for corrida in parrafo_celda.runs:
                corrida.font.bold = True
    for fila in filas:
        celdas = tabla.add_row().cells
        for indice, (celda, texto) in enumerate(zip(celdas, fila)):
            celda.text = str(texto)
            for parrafo_celda in celda.paragraphs:
                parrafo_celda.alignment = (WD_ALIGN_PARAGRAPH.LEFT if indice in izquierda
                                           else WD_ALIGN_PARAGRAPH.CENTER)
    for fila in tabla.rows:
        for celda in fila.cells:
            for parrafo_celda in celda.paragraphs:
                parrafo_celda.paragraph_format.line_spacing = 1.0
                parrafo_celda.paragraph_format.space_after = Pt(2)
                for corrida in parrafo_celda.runs:
                    corrida.font.name = "Arial"
                    corrida.font.size = Pt(9)
    doc.add_paragraph()
    return tabla


def agrupar(cuadro, pares, columna_inicial=0):
    for izquierda, derecha in pares:
        cuadro.cell(0, izquierda).merge(cuadro.cell(0, derecha))
    for columna in range(columna_inicial + 1):
        cuadro.cell(0, columna).merge(cuadro.cell(1, columna))


def portada(doc):
    def centrado(texto, tamano=12, negrita=False, espacio=6):
        parrafo_actual = doc.add_paragraph()
        parrafo_actual.alignment = WD_ALIGN_PARAGRAPH.CENTER
        parrafo_actual.paragraph_format.space_after = Pt(espacio)
        corrida = parrafo_actual.add_run(texto)
        corrida.font.size = Pt(tamano)
        corrida.font.bold = negrita
        return parrafo_actual

    centrado("UNIVERSIDAD DE INGENIERÍA Y TECNOLOGÍA", 14, True, 24)
    doc.add_picture(LOGO, width=Cm(6))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(24)
    centrado("Curso:", 14, True, 2)
    centrado("Base de Datos II", 14, False, 36)
    centrado("Laboratorio 04: Sequential File vs AVL File", 18, True, 48)
    centrado("Alumno:", 12, True, 2)
    centrado("Llerena Silva, Nicolás Alejandro (202110190)", 12, False, 24)
    centrado("Profesor:", 12, True, 2)
    centrado("Lovon, Percy", 12, False, 36)
    centrado("Lima, setiembre de 2026", 12, False, 0)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def introduccion(doc):
    doc.add_heading("Introducción", level=1)
    parrafo(doc, "El laboratorio pide implementar dos organizaciones de archivo sobre memoria "
                 "secundaria y compararlas usando el Employee_ID como clave: un archivo secuencial "
                 "ordenado con espacio auxiliar y un archivo organizado como árbol AVL con los "
                 "punteros guardados en el propio archivo. Sobre cada uno se implementaron insert, "
                 "search, remove y rangeSearch, y después se midió cuánto cuesta cada operación al "
                 "crecer el archivo.")
    parrafo(doc, "Todo está escrito en Python con la librería estándar; solo el script de gráficas y "
                 "el que arma este documento usan matplotlib y python-docx. Los archivos son binarios "
                 "de longitud fija y las mediciones se hicieron en una MacBook con SSD, sobre el mismo "
                 "conjunto de datos para las dos estructuras.")

    doc.add_heading("Los datos y el registro", level=2)
    parrafo(doc, "El enunciado parte de un employee.csv que no venía con el material, así que lo "
                 "construí desde la base employees del Laboratorio 01, que ya tiene empleados reales "
                 "con nombre, departamento, puesto, sueldo y fecha de ingreso. El script "
                 "exportar_csv.py toma los 20 500 empleados vigentes (los que tienen departamento, "
                 "título y salario con to_date = 9999-01-01) y arma el CSV con las ocho columnas que "
                 "pide el enunciado. El único campo que la base no tiene es Country, así que se asigna "
                 "de forma determinista a partir del id; los otros siete son datos reales. Si más "
                 "adelante aparece el employee.csv oficial, los dos programas lo cargan igual: el "
                 "lector normaliza los nombres de columna y no depende del orden en que vengan.")
    parrafo(doc, "El registro sigue exactamente la tabla del enunciado y ocupa 112 bytes. Las cadenas "
                 "se guardan rellenadas con espacios hasta su tamaño máximo y se recortan al leer. "
                 "Vale la pena anotar un detalle del Salary: al pedirse en 4 bytes es un float de "
                 "precisión simple, que tiene unos siete dígitos significativos, así que sueldos como "
                 "95000 se guardan exactos pero un valor con muchos decimales se redondea al leerlo.")
    tabla(doc,
          ["Campo", "Tipo", "Bytes", "Desplazamiento"],
          [["Employee_ID", "int", "4", "0"],
           ["Employee_Name", "string", "30", "4"],
           ["Age", "int", "4", "34"],
           ["Country", "string", "20", "38"],
           ["Department", "string", "20", "58"],
           ["Position", "string", "20", "78"],
           ["Salary", "float", "4", "98"],
           ["Joining_Date", "string DD/MM/YYYY", "10", "102"],
           ["Total", "", "112", ""]],
          izquierda=(0, 1))
    captura(doc, "cap_datos.png", "Exportación del CSV desde la base employees del Laboratorio 01.")


def parte1(doc):
    doc.add_heading("P1: archivo secuencial", level=1)
    parrafo(doc, "El archivo secuencial son dos archivos: el de datos, ordenado por Employee_ID, y "
                 "uno auxiliar donde caen las inserciones sin orden. Cada registro se guarda con un "
                 "byte extra de estado al final, así que ocupa 113 bytes en disco: 112 de datos y uno "
                 "que vale 1 si el registro está activo y 0 si fue eliminado. No hay cabecera porque "
                 "la cantidad de registros sale del tamaño del archivo dividido entre 113.")
    parrafo(doc, "insert escribe siempre al final del auxiliar, que cuesta un solo acceso. Cuando el "
                 "auxiliar supera k registros se dispara la reconstrucción: se lee el archivo de datos "
                 "de principio a fin, se mezcla con el auxiliar ya ordenado en memoria (nunca son más "
                 "de k registros) y se escribe un archivo nuevo que reemplaza al anterior. En esa "
                 "mezcla se descartan los registros marcados como eliminados, que es donde el borrado "
                 "lógico se hace efectivo. La mezcla es en streaming, o sea que el archivo de datos "
                 "nunca se carga entero en memoria.")
    parrafo(doc, "search hace búsqueda binaria sobre el archivo de datos, que son log₂(n) accesos, y "
                 "si no encuentra la clave recorre el auxiliar, que son a lo más k accesos más. remove "
                 "usa la misma búsqueda y se limita a poner el byte de estado en 0, sin mover nada de "
                 "sitio. rangeSearch busca con búsqueda binaria la primera clave mayor o igual al "
                 "inicio del rango y desde ahí avanza secuencialmente mientras la clave siga dentro "
                 "del rango, que es la parte donde este diseño se luce: los registros que necesita "
                 "están uno al lado del otro en el archivo.")
    captura(doc, "cap_seq.png", "Prueba de las cuatro operaciones sobre 2000 empleados, con k = 16. "
                                "Entre corchetes van los accesos a disco de cada operación.")


def parte2(doc):
    doc.add_heading("P2: archivo AVL", level=1)
    parrafo(doc, "En el archivo AVL cada nodo es un registro de 112 bytes seguido de tres enteros: la "
                 "posición del hijo izquierdo, la del derecho y la altura del subárbol. Son 124 bytes "
                 "por nodo y las posiciones son índices dentro del archivo, no offsets en bytes, con "
                 "−1 como puntero nulo. Al inicio del archivo hay una cabecera de 12 bytes con la "
                 "posición de la raíz, la cabeza de la lista de posiciones libres y la cantidad de "
                 "nodos vivos.")
    parrafo(doc, "insert baja por el árbol comparando el Employee_ID, crea el nodo en la primera "
                 "posición libre que encuentre y al volver de la recursión recalcula alturas y "
                 "rebalancea. Las rotaciones son las cuatro clásicas y cada una reescribe los dos "
                 "nodos que cambian de lugar más los punteros del padre, todo sobre el archivo. remove "
                 "es la eliminación estándar del AVL: si el nodo tiene un hijo o ninguno se enlaza el "
                 "hijo al padre y la posición liberada se encadena en la lista de libres para "
                 "reutilizarla en la siguiente inserción; si tiene dos hijos se copia el registro del "
                 "sucesor inorden y se elimina ese sucesor del subárbol derecho. En ambos casos se "
                 "rebalancea al regresar.")
    parrafo(doc, "search es un descenso iterativo, un acceso por nivel. rangeSearch es un recorrido "
                 "inorden podado: si la clave del nodo ya es menor o igual al inicio del rango no se "
                 "baja por la izquierda, y si es mayor o igual al final no se baja por la derecha, así "
                 "que solo se visitan los nodos del rango más el camino que lleva a ellos. Como el "
                 "recorrido inorden de un árbol de búsqueda sale ordenado, el resultado ya viene "
                 "ordenado por Employee_ID sin necesidad de ordenarlo después.")
    captura(doc, "cap_avl.png", "Las mismas pruebas sobre el AVL. La altura 11 con 2000 nodos "
                                "confirma que el árbol quedó balanceado, y el recorrido inorden "
                                "devuelve los registros en orden.")


def parte3(doc, datos):
    filas = datos["tamanos"]
    factores = datos["factor_k"]
    grande = por_tamano(filas, 20000)
    chico = por_tamano(filas, 1000)

    doc.add_heading("P3: evaluación de desempeño", level=1)
    parrafo(doc, "Para medir usé cinco tamaños de archivo (1000, 2500, 5000, 10 000 y 20 000 "
                 "registros) tomados del mismo CSV barajado con una semilla fija, y 500 registros "
                 "apartados que nunca entran en la construcción y sirven para las pruebas de "
                 "inserción. Sobre cada tamaño se ejecutan 500 búsquedas de claves existentes, 100 "
                 "búsquedas por rango de 250 unidades de ancho, las 500 inserciones y 500 "
                 "eliminaciones, siempre con las mismas claves para las dos estructuras. Además de "
                 "cronometrar cada operación cuento los accesos a disco, que son las lecturas y "
                 "escrituras de registro o de nodo que hace cada estructura; esa métrica es la que "
                 "de verdad describe el costo en memoria secundaria, porque el tiempo se contamina con "
                 "la caché del sistema operativo.")
    parrafo(doc, "El benchmark también compara los resultados de las dos estructuras: cada búsqueda "
                 "y cada rango tienen que devolver lo mismo en el archivo secuencial y en el AVL, y "
                 "al final los dos deben contener el mismo conjunto de registros válidos. Esas "
                 "comprobaciones son asserts dentro del programa, así que si algo no coincidiera el "
                 "benchmark se caería en vez de reportar números.")

    cuadro = tabla(doc,
          ["Registros", "Inserción de 500 (s)", "", "Búsqueda (ms)", "", "Rango (ms)", "",
           "Eliminación (ms)", ""],
          [["", "Sec.", "AVL", "Sec.", "AVL", "Sec.", "AVL", "Sec.", "AVL"]] +
          [[fila["tamano"],
            coma(fila["insercion"]["secuencial"]["total_seg"]),
            coma(fila["insercion"]["avl"]["total_seg"]),
            coma(fila["busqueda"]["secuencial"]["promedio_ms"], 3),
            coma(fila["busqueda"]["avl"]["promedio_ms"], 3),
            coma(fila["rango"]["secuencial"]["promedio_ms"], 3),
            coma(fila["rango"]["avl"]["promedio_ms"], 3),
            coma(fila["eliminacion"]["secuencial"]["promedio_ms"], 3),
            coma(fila["eliminacion"]["avl"]["promedio_ms"], 3)] for fila in filas])
    agrupar(cuadro, [(1, 2), (3, 4), (5, 6), (7, 8)], columna_inicial=0)
    doc.add_paragraph("Tiempos por operación. La inserción es el total de las 500; las demás son el "
                      "promedio por operación.", style="Pie")

    cuadro = tabla(doc,
          ["Registros", "Altura AVL", "Inserción", "", "Búsqueda", "", "Rango", "",
           "Eliminación", ""],
          [["", "", "Sec.", "AVL", "Sec.", "AVL", "Sec.", "AVL", "Sec.", "AVL"]] +
          [[fila["tamano"], fila["altura_avl"],
            f"{fila['insercion']['secuencial']['accesos_promedio']:.0f}",
            f"{fila['insercion']['avl']['accesos_promedio']:.0f}",
            f"{fila['busqueda']['secuencial']['accesos_promedio']:.1f}",
            f"{fila['busqueda']['avl']['accesos_promedio']:.1f}",
            f"{fila['rango']['secuencial']['accesos_promedio']:.0f}",
            f"{fila['rango']['avl']['accesos_promedio']:.0f}",
            f"{fila['eliminacion']['secuencial']['accesos_promedio']:.0f}",
            f"{fila['eliminacion']['avl']['accesos_promedio']:.0f}"] for fila in filas])
    agrupar(cuadro, [(2, 3), (4, 5), (6, 7), (8, 9)], columna_inicial=1)
    doc.add_paragraph("Accesos a disco promedio por operación.", style="Pie")
    captura(doc, "cap_bench.png", "Salida del benchmark completo.")
    figura(doc, os.path.join(RESULTADOS, "tiempos.png"),
           "Tiempos de las cuatro operaciones al crecer el archivo.", ancho=15.0)
    figura(doc, os.path.join(RESULTADOS, "accesos.png"),
           "Accesos a disco por operación con 20 000 registros, en escala logarítmica.")

    doc.add_heading("Análisis por operación", level=2)
    parrafo(doc, f"<b>Inserción.</b> Es la diferencia más grande y la que decide la comparación. El "
                 f"AVL se mantiene plano: pasa de {chico['insercion']['avl']['accesos_promedio']:.0f} "
                 f"a {grande['insercion']['avl']['accesos_promedio']:.0f} accesos por inserción entre "
                 f"1000 y 20 000 registros, un crecimiento logarítmico, y las 500 inserciones tardan "
                 f"{coma(grande['insercion']['avl']['total_seg'])} s en el archivo más grande. El "
                 f"secuencial va de {chico['insercion']['secuencial']['accesos_promedio']:.0f} a "
                 f"{grande['insercion']['secuencial']['accesos_promedio']:.0f} accesos, o sea que "
                 f"crece de forma lineal con el tamaño del archivo, y termina tardando "
                 f"{coma(grande['insercion']['secuencial']['total_seg'])} s, "
                 f"{coma(grande['insercion']['secuencial']['total_seg'] / grande['insercion']['avl']['total_seg'], 1)} "
                 f"veces más. La inserción en sí es barata, un acceso al auxiliar; lo caro son las "
                 f"{grande['reconstrucciones']} reconstrucciones que se disparan, cada una leyendo y "
                 f"reescribiendo el archivo completo.")
    parrafo(doc, f"<b>Búsqueda.</b> Aquí quedan parejos, y tiene sentido: los dos son logarítmicos. "
                 f"Con 20 000 registros la búsqueda binaria usa "
                 f"{grande['busqueda']['secuencial']['accesos_promedio']:.1f} accesos y el AVL "
                 f"{grande['busqueda']['avl']['accesos_promedio']:.1f}, con el árbol apenas por encima "
                 f"porque su altura ({grande['altura_avl']}) es mayor que log₂(20 000) ≈ 14,3; un AVL "
                 f"puede llegar a 1,44 veces esa cota. En tiempo el AVL sale levemente mejor "
                 f"({coma(grande['busqueda']['avl']['promedio_ms'], 3)} ms contra "
                 f"{coma(grande['busqueda']['secuencial']['promedio_ms'], 3)} ms) porque la búsqueda "
                 f"binaria salta a posiciones muy separadas del archivo mientras que los nodos altos "
                 f"del árbol se visitan en todas las búsquedas y quedan en caché. La diferencia es de "
                 f"microsegundos: para búsquedas puntuales las dos estructuras son equivalentes.")
    parrafo(doc, f"<b>Búsqueda por rango.</b> Se invierte la ventaja. Los accesos son casi idénticos "
                 f"({grande['rango']['secuencial']['accesos_promedio']:.0f} contra "
                 f"{grande['rango']['avl']['accesos_promedio']:.0f} con 20 000 registros, porque los "
                 f"dos terminan tocando los ~{grande['registros_por_rango']:.0f} registros del rango), "
                 f"pero el tiempo del secuencial es la mitad: "
                 f"{coma(grande['rango']['secuencial']['promedio_ms'], 3)} ms contra "
                 f"{coma(grande['rango']['avl']['promedio_ms'], 3)} ms. La razón es la localidad: el "
                 f"secuencial encuentra el inicio del rango y de ahí lee registros contiguos, que el "
                 f"sistema operativo trae de a bloques enteros, mientras que el AVL recorre el árbol "
                 f"saltando entre nodos dispersos por todo el archivo. Es exactamente el argumento por "
                 f"el que un índice clustered le gana a uno no clustered en consultas de rango.")
    parrafo(doc, f"<b>Eliminación.</b> Gana el secuencial y con holgura: "
                 f"{grande['eliminacion']['secuencial']['accesos_promedio']:.0f} accesos contra "
                 f"{grande['eliminacion']['avl']['accesos_promedio']:.0f} del AVL. Pero la comparación "
                 f"es un poco tramposa, porque no hacen lo mismo. El secuencial hace borrado lógico: "
                 f"busca el registro y le cambia un byte, así que cuesta lo mismo que una búsqueda más "
                 f"una escritura, y el espacio sigue ocupado hasta la próxima reconstrucción. El AVL "
                 f"elimina de verdad: reestructura el árbol, puede rotar en todo el camino de vuelta a "
                 f"la raíz y devuelve la posición liberada a la lista de libres para reutilizarla. El "
                 f"secuencial paga después lo que no paga ahora.")

    doc.add_heading("El parámetro k del archivo secuencial", level=2)
    parrafo(doc, f"El tamaño del auxiliar es la perilla que ajusta el archivo secuencial, así que "
                 f"medí aparte cómo se comporta con 10 000 registros. Con k = {factores[0]['k']} las "
                 f"500 inserciones provocan {factores[0]['reconstrucciones']} reconstrucciones y "
                 f"tardan {coma(factores[0]['insercion_total_seg'])} s; con k = {factores[-1]['k']} "
                 f"bajan a {factores[-1]['reconstrucciones']} reconstrucciones y "
                 f"{coma(factores[-1]['insercion_total_seg'])} s, catorce veces más rápido. El precio "
                 f"aparece del otro lado: buscar un registro que todavía está en el auxiliar cuesta "
                 f"{factores[0]['busqueda_auxiliar_accesos']:.0f} accesos con k = {factores[0]['k']} y "
                 f"{factores[-1]['busqueda_auxiliar_accesos']:.0f} con k = {factores[-1]['k']}, "
                 f"porque el auxiliar se recorre de forma lineal. Las búsquedas que caen en el archivo "
                 f"de datos no se enteran: se mantienen en unos "
                 f"{factores[0]['busqueda_accesos']:.0f} accesos en todos los casos.")
    figura(doc, os.path.join(RESULTADOS, "factor_k.png"),
           "Efecto de k sobre el costo de insertar y sobre el de buscar en el auxiliar.")

    doc.add_heading("Cuándo conviene cada uno", level=2)
    parrafo(doc, "El archivo secuencial conviene cuando los datos casi no cambian y lo que se hace "
                 "sobre ellos son lecturas, sobre todo por rango: reportes, históricos, tablas de "
                 "referencia, cualquier cosa que se cargue una vez y después solo se consulte. Su "
                 "construcción es prácticamente gratis comparada con la del AVL (0,02 s contra 2,4 s "
                 "para 20 000 registros, porque una es ordenar y volcar y la otra son 20 000 "
                 "inserciones con rotaciones), la búsqueda binaria es óptima en accesos y el barrido "
                 "por rango aprovecha la localidad del disco. También es más simple de implementar y "
                 "de recuperar ante una falla.")
    parrafo(doc, "El AVL conviene cuando el archivo se modifica seguido y no se puede pagar una "
                 "reconstrucción cada tantas inserciones: catálogos que crecen todo el día, tablas "
                 "transaccionales, cualquier caso donde insertar y borrar sean tan frecuentes como "
                 "consultar. Mantiene el costo logarítmico sin importar el orden en que lleguen las "
                 "claves y su eliminación libera espacio de inmediato. A cambio, cada nodo carga 12 "
                 "bytes de punteros (un 11 % de sobrecarga sobre el registro) y el recorrido por rango "
                 "pierde localidad.")
    parrafo(doc, "Dicho de otra forma: si el archivo se lee mucho más de lo que se escribe, gana el "
                 "secuencial; si se escribe seguido, gana el AVL. Y si hicieran falta las dos cosas a "
                 "la vez, el camino no sería elegir uno sino usar un árbol B+, que junta la altura "
                 "logarítmica del AVL con las hojas enlazadas que dan localidad en los rangos.")


def conclusiones(doc):
    doc.add_heading("Conclusiones", level=1)
    parrafo(doc, "Las dos estructuras resuelven las cuatro operaciones sobre los mismos datos y "
                 "devuelven exactamente los mismos resultados, así que la comparación es sobre el "
                 "costo y no sobre la funcionalidad. La búsqueda puntual es un empate, porque las dos "
                 "son logarítmicas. En inserción el AVL gana claramente y la distancia crece con el "
                 "archivo, ya que el secuencial reconstruye en tiempo lineal. En rango gana el "
                 "secuencial por localidad, con la mitad del tiempo pese a hacer los mismos accesos. "
                 "Y en eliminación gana el secuencial solo porque no hace el trabajo completo: marca "
                 "el registro y difiere el costo real a la siguiente reconstrucción.")
    parrafo(doc, "La medición también deja ver que contar accesos a disco y cronometrar no son lo "
                 "mismo. En el rango los accesos son iguales y los tiempos difieren al doble, y en la "
                 "búsqueda pasa al revés. Cuando el archivo entra en la caché del sistema operativo, "
                 "como acá, lo que termina pesando es el patrón de acceso más que la cantidad.")


def anexo(doc):
    doc.add_heading("Anexo: archivos entregados", level=1)
    parrafo(doc, "registro.py define el registro de 112 bytes, su empaquetado con struct y la lectura "
                 "del CSV. sequential_file.py implementa la clase SequentialFile con build, insert, "
                 "search, remove, range_search y la reconstrucción por streaming. avl_file.py "
                 "implementa AVLFile con las mismas operaciones sobre el árbol en disco, incluyendo "
                 "rotaciones y lista de posiciones libres. exportar_csv.py arma data/employee.csv "
                 "desde PostgreSQL y benchmark.py corre las mediciones y genera las gráficas. Los tres "
                 "módulos principales se ejecutan solos y vuelven a imprimir las salidas que aparecen "
                 "en las capturas de este informe.")


def main():
    datos = metricas()
    doc = documento()
    portada(doc)
    introduccion(doc)
    parte1(doc)
    parte2(doc)
    parte3(doc, datos)
    conclusiones(doc)
    anexo(doc)
    doc.save(SALIDA)
    print(f"{SALIDA} ({os.path.getsize(SALIDA) // 1024} KB)")


if __name__ == "__main__":
    main()
